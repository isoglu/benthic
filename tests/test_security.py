"""Schutz der lokalen Schnittstelle (Host, Origin, Art der Anfrage) und verdeckte API-Keys (offline)."""

import json
import os
import stat
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
from pathlib import Path
from unittest import mock

from tests import ui_client


class Regel(unittest.TestCase):
    """app.refusal: wer darf was?"""

    def setUp(self):
        from benthic import app as appmod
        self.refusal = appmod.refusal

    def check(self, *, host="127.0.0.1", port="8327", method="GET", origin=None, ctype="",
              site="", mode="", own="8327"):
        return self.refusal(host, port, method, origin, ctype, site, mode, own)

    def test_host(self):
        self.assertIsNone(self.check())
        self.assertIsNone(self.check(host="localhost"))
        self.assertEqual(self.check(host="evil.example"), "host")        # DNS-Rebinding
        self.assertEqual(self.check(host="127.0.0.1.evil.example"), "host")
        self.assertEqual(self.check(port="8328"), "host")                # anderer Port als der eigene
        self.assertEqual(self.check(host="[::1]:8327", port="80"), "host")

    def test_origin(self):
        own = "http://127.0.0.1:8327"
        self.assertIsNone(self.check(origin=own))
        self.assertIsNone(self.check(origin="http://localhost:8327", method="POST"))
        self.assertEqual(self.check(origin="https://evil.example"), "origin")
        self.assertEqual(self.check(origin="https://evil.example", method="POST", ctype="application/json"), "origin")
        self.assertEqual(self.check(origin="http://127.0.0.1:9999", method="POST"), "origin")
        self.assertEqual(self.check(origin="null", method="POST", ctype="text/plain"), "origin")
        self.assertIsNone(self.check(origin="null", method="POST", ctype="application/json"))  # etwa Zotero

    def test_state_changing_requests(self):
        # MCP-Server und Zotero-Plug-in: JSON ohne Origin
        self.assertIsNone(self.check(method="POST", ctype="application/json"))
        self.assertIsNone(self.check(method="DELETE", ctype="application/json"))
        # Einfache Anfragen, die jede Website ohne Vorabprüfung schicken kann
        for ctype in ("text/plain", "application/x-www-form-urlencoded", "multipart/form-data", ""):
            with self.subTest(ctype=ctype):
                self.assertEqual(self.check(method="POST", ctype=ctype), "type")
        # Die eigene Oberfläche: Uploads und Aufrufe ohne Inhalt
        self.assertIsNone(self.check(method="POST", ctype="multipart/form-data", origin="http://127.0.0.1:8327"))
        self.assertIsNone(self.check(method="POST", origin="http://127.0.0.1:8327"))
        self.assertEqual(self.check(method="OPTIONS", origin="https://evil.example"), "origin")

    def test_fetch_metadata(self):
        self.assertEqual(self.check(site="cross-site", mode="no-cors"), "site")  # <img>, <script> fremder Seiten
        self.assertEqual(self.check(site="same-site", mode="cors"), "site")
        self.assertIsNone(self.check(site="cross-site", mode="navigate"))       # Link auf einen Bericht
        self.assertIsNone(self.check(site="same-origin", mode="cors"))
        self.assertIsNone(self.check(site="none", mode="navigate"))


class Schnittstelle(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.env = mock.patch.dict(os.environ, {"BENTHIC_DATA_DIR": self.tmp.name})
        self.env.start()
        from benthic import app as appmod, config
        self.config = config
        self.app = appmod.app
        self.client = self.app.test_client()

    def tearDown(self):
        self.env.stop()
        self.tmp.cleanup()

    def test_foreign_host_is_refused(self):
        self.assertEqual(self.client.get("/api/health").status_code, 200)
        r = self.client.get("/api/settings", headers={"Host": "evil.example"})
        self.assertEqual(r.status_code, 403)
        self.assertNotIn("api_key", r.get_data(as_text=True))

    def test_simple_cross_site_posts_change_nothing(self):
        from benthic import library
        with mock.patch.object(library, "clear", return_value=5) as clear:
            r = self.client.post("/api/library/clear", data="x", content_type="text/plain")
            self.assertEqual(r.status_code, 403)
            r = self.client.post("/api/library/clear", headers={"Origin": "https://evil.example"})
            self.assertEqual(r.status_code, 403)
            clear.assert_not_called()
            self.assertEqual(ui_client(self.app).post("/api/library/clear").status_code, 200)
            clear.assert_called_once()
        r = self.client.post("/api/projects/import", data=json.dumps({"name": "x", "items": []}), content_type="text/plain")
        self.assertEqual(r.status_code, 403)
        self.assertEqual(self.client.get("/api/projects").get_json()["projects"], [])

    def test_json_without_origin_still_works(self):
        """So spricht der MCP-Server mit dem Backend."""
        r = self.client.post("/api/projects", json={"name": "Seminar"})
        self.assertEqual(r.status_code, 200, r.get_json())

    def test_api_keys_are_masked(self):
        key = "abcdefghijklmnopWXYZ"
        self.client.post("/api/settings", json={"core_api_key": key, "elsevier_api_key": "short"})
        shown = self.client.get("/api/settings").get_json()
        self.assertEqual(shown["core_api_key"], self.config.MASK + "WXYZ")
        self.assertEqual(shown["elsevier_api_key"], self.config.MASK)
        self.assertEqual(shown["openalex_api_key"], "")
        self.assertNotIn(key, json.dumps(shown))
        # Die Oberfläche speichert alle Felder: unveränderte Platzhalter lassen die Keys stehen.
        back = self.client.post("/api/settings", json=dict(shown, concurrency=5)).get_json()
        self.assertEqual(back["core_api_key"], self.config.MASK + "WXYZ")
        saved = self.config.load_settings()
        self.assertEqual((saved["core_api_key"], saved["elsevier_api_key"], saved["concurrency"]), (key, "short", 5))
        # Leeren löscht, Neues ersetzt.
        self.client.post("/api/settings", json={"core_api_key": "", "elsevier_api_key": "new-key-1234567"})
        saved = self.config.load_settings()
        self.assertEqual((saved["core_api_key"], saved["elsevier_api_key"]), ("", "new-key-1234567"))
        mode = stat.S_IMODE(self.config.settings_path().stat().st_mode)
        self.assertEqual(mode, 0o600)

    def test_setup_answer_masks_keys(self):
        self.config.save_settings({"openalex_api_key": "secret-openalex-key"})
        r = self.client.post("/api/setup", json={"email": "a@b.example", "download_dir": self.tmp.name})
        self.assertEqual(r.status_code, 200, r.get_json())
        self.assertNotIn("secret-openalex-key", r.get_data(as_text=True))

    def test_request_size_is_capped(self):
        import io
        from benthic import document
        self.assertLessEqual(self.app.config["MAX_CONTENT_LENGTH"], document.MAX_BYTES + 8 * 1024 * 1024)
        with mock.patch.dict(self.app.config, {"MAX_CONTENT_LENGTH": 1000}):
            r = ui_client(self.app).post("/api/document", data={"file": (io.BytesIO(b"x" * 5000), "a.txt")},
                                         content_type="multipart/form-data")
        self.assertEqual(r.status_code, 413)
        self.assertIn("60 MB", r.get_json()["error"])


class EchterServer(unittest.TestCase):
    """Der Werkzeug-Server auf einem freien Port: Host mit Port, MCP-Client, fremde Hosts."""

    def setUp(self):
        from werkzeug.serving import make_server
        self.tmp = tempfile.TemporaryDirectory()
        self.env = mock.patch.dict(os.environ, {"BENTHIC_DATA_DIR": self.tmp.name})
        self.env.start()
        from benthic import app as appmod
        self.server = make_server("127.0.0.1", 0, appmod.app, threaded=True)
        self.port = self.server.server_port
        threading.Thread(target=self.server.serve_forever, daemon=True).start()

    def tearDown(self):
        self.server.shutdown()
        self.server.server_close()
        self.env.stop()
        self.tmp.cleanup()

    def get(self, path, headers=None):
        req = urllib.request.Request(f"http://127.0.0.1:{self.port}{path}", headers=headers or {})
        try:
            with urllib.request.urlopen(req, timeout=5) as resp:
                return resp.status
        except urllib.error.HTTPError as exc:
            return exc.code

    def test_hosts(self):
        self.assertEqual(self.get("/api/health"), 200)
        self.assertEqual(self.get("/api/health", {"Host": f"localhost:{self.port}"}), 200)
        self.assertEqual(self.get("/api/health", {"Host": f"rebind.example:{self.port}"}), 403)
        self.assertEqual(self.get("/api/health", {"Host": "127.0.0.1:1"}), 403)
        self.assertEqual(self.get("/", {"Origin": "https://evil.example"}), 403)

    def test_mcp_client(self):
        from benthic import mcp_server
        with mock.patch.object(mcp_server, "BASE", f"http://127.0.0.1:{self.port}"):
            self.assertEqual(mcp_server._http("GET", "/api/health")[0], 200)
            status, body = mcp_server._http("POST", "/api/projects", {"name": "MCP"})
            self.assertEqual(status, 200, body)
            pid = body["id"]
            self.assertEqual(mcp_server._http("DELETE", f"/api/projects/{pid}")[0], 200)


class Dateinamen(unittest.TestCase):
    """Namen geladener PDFs kommen aus fremden Metadaten und bleiben im Zielordner."""

    def test_every_part_is_slugified(self):
        from benthic.crossref import build_filename
        name = build_filename("10.1/x", {"author": "../../Library/LaunchAgents/planted", "year": "2024/..",
                                         "title": "A study"})
        self.assertRegex(name, r"^[A-Za-z0-9_]+\.pdf$")
        self.assertEqual(build_filename("10.1/x", {"author": "Smith", "year": "2020", "title": "A study"}),
                         "Smith_2020_A_study.pdf")
        self.assertEqual(build_filename("10.1/x", {"author": "Müller", "year": 2019, "title": ""}), "Muller_2019.pdf")
        self.assertEqual(build_filename("10.1234/ab.c", {"author": "/", "year": "", "title": "//"}), "10_1234_ab_c.pdf")
        self.assertEqual(build_filename("///", None), "paper.pdf")

    def test_save_stays_in_the_folder(self):
        from benthic import crossref, fetcher
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "papers"
            d = fetcher.BenthicDownloader(out, skip_existing=False)
            with mock.patch.object(crossref, "build_filename", return_value="../escape.pdf"):
                self.assertIsNone(d._save(b"%PDF", "10.1/x", {}))
            self.assertFalse((Path(tmp) / "escape.pdf").exists())
            self.assertEqual(d._save(b"%PDF", "10.1/x", {"author": "../x", "year": "1", "title": "t"}), "x_1_t.pdf")
            self.assertEqual([p.name for p in out.iterdir()], ["x_1_t.pdf"])


def _zip(members):
    import io
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in members.items():
            z.writestr(name, data)
    return buf.getvalue()


class Groessen(unittest.TestCase):
    """Entpacken und Lesen mit Obergrenzen: Archive, Word, Pfade, Ordner."""

    def test_zip_total_is_capped(self):
        from benthic import document, latex
        data = _zip({f"ch{i}.tex": "x" * 600 for i in range(3)})
        self.assertEqual(len(latex.files_from_upload([("p.zip", data)])), 3)
        with mock.patch.object(document, "MAX_UNPACKED", 1000):
            with self.assertRaises(ValueError) as cm:
                latex.files_from_upload([("p.zip", data)])
        self.assertEqual(str(cm.exception), "size")

    def test_docx_total_is_capped(self):
        from benthic import document
        body = '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>' \
               + "<w:p><w:r><w:t>Text</w:t></w:r></w:p>" * 100 + "</w:body></w:document>"
        data = _zip({"word/document.xml": body})
        self.assertIn("Text", document.extract_text("a.docx", data))
        with mock.patch.object(document, "MAX_UNPACKED", 1000):
            with self.assertRaises(ValueError) as cm:
                document.extract_text("a.docx", data)
        self.assertEqual(str(cm.exception), "size")

    def test_xlsx_total_is_capped(self):
        from benthic import document, journals
        with mock.patch.object(document, "MAX_UNPACKED", 100):
            with self.assertRaises(ValueError):
                journals.read_table("r.xlsx", _zip({"xl/workbook.xml": "x" * 500}))

    def test_paths_are_checked_before_reading(self):
        from benthic import app as appmod
        with tempfile.TemporaryDirectory() as tmp:
            fifo = Path(tmp) / "thesis.pdf"
            os.mkfifo(fifo)  # Lesen würde ewig warten
            big = Path(tmp) / "big.pdf"
            with big.open("wb") as f:
                f.truncate(61 * 1024 * 1024)  # dünn besetzt, belegt keinen Platz
            exe = Path(tmp) / "tool.exe"
            exe.write_bytes(b"MZ")
            for path, reason in ((fifo, "type"), (Path("/dev/zero"), "type"), (big, "size"), (exe, "type")):
                with self.subTest(path=path.name):
                    with self.assertRaises(ValueError) as cm:
                        appmod._analyse_document([(path.name, b"")], path)
                    self.assertEqual(str(cm.exception), reason)
            client = appmod.app.test_client()
            for path in (fifo, "/dev/zero"):
                r = client.post("/api/document", json={"path": str(path)})
                self.assertEqual(r.status_code, 400)

    def test_folder_scan_is_bounded(self):
        from benthic import latex
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "main.tex").write_text("\\documentclass{article}")
            (root / ".git").mkdir()
            (root / ".git" / "x.tex").write_text("hidden")
            for i in range(8):
                (root / f"img{i}.png").write_bytes(b"x")
            self.assertEqual(sorted(latex.files_from_path(root)), ["main.tex"])
            with mock.patch.object(latex, "MAX_SCAN", 5):
                with self.assertRaises(ValueError) as cm:
                    latex.files_from_path(root)
            self.assertEqual(str(cm.exception), "too-many")

ROOT = Path(__file__).resolve().parent.parent


class Links(unittest.TestCase):
    """Nur Web-Links aus fremden Metadaten; die App gibt nur bekannte Schemata an macOS weiter."""

    def test_report_links(self):
        from benthic import report
        for bad in ("javascript:alert(1)", "file:///etc/passwd", "smb://host/share", " vnc://x", "", None):
            self.assertEqual(report._web(bad), "", bad)
        self.assertEqual(report._web("https://web.archive.org/web/2020/x"), "https://web.archive.org/web/2020/x")
        card = report._attention_card({"verdict": "web", "raw": "Seite", "status": "web",
                                       "web": {"url": "https://example.org/a", "archive_url": "file:///etc/passwd"}})
        self.assertNotIn("file:", card)
        card = report._attention_card({"verdict": "web", "raw": "Seite", "status": "web",
                                       "web": {"url": "https://example.org/a",
                                               "archive_url": "https://web.archive.org/web/2020/https://example.org/a"}})
        self.assertIn('href="https://web.archive.org/web/2020/https://example.org/a"', card)

    def test_app_js_filters_metadata_links(self):
        src = (ROOT / "benthic" / "static" / "app.js").read_text(encoding="utf-8")
        for raw in ("href = res.landing_url", "href = j.landing_url", "href = c.orcid", "href = w.oa_url",
                    "link(web.archive_url"):
            self.assertNotIn(raw, src)
        import shutil
        import subprocess
        if not shutil.which("node"):
            self.skipTest("node fehlt")
        fn = next(l for l in src.splitlines() if l.startswith("function webUrl("))
        code = fn + "\nconsole.log(JSON.stringify(['https://a.example/x', 'HTTP://b.example', 'javascript:alert(1)'," \
                    " 'file:///etc/passwd', ' vnc://x', null, 'zotero://select/x'].map(webUrl)));"
        out = subprocess.run(["node", "-e", code], capture_output=True, text=True, check=True).stdout
        self.assertEqual(json.loads(out), ["https://a.example/x", "HTTP://b.example", "", "", "", "", ""])

    def test_shell_hands_only_known_schemes_to_macos(self):
        swift = (ROOT / "macos" / "BenthicApp.swift").read_text(encoding="utf-8")
        self.assertNotIn("hasPrefix(backendHost", swift)  # Präfixvergleich ließe http://127.0.0.1:8327@fremd durch
        self.assertIn('let externalSchemes: Set<String> = ["http", "https", "mailto", "zotero", "vscode", "cursor", '
                      '"lmstudio"]', swift)
        opens = [l.strip() for l in swift.splitlines() if "NSWorkspace.shared.open(" in l]
        self.assertEqual(opens, ["NSWorkspace.shared.open(url)",                      # openExternally
                                 "if ok { NSWorkspace.shared.open(dmg) }",            # Update erst nach der Signatur
                                 "NSWorkspace.shared.open(url)"])                     # Lizenzen im Bundle
        self.assertEqual(swift.count("openExternally("), 4)  # Definition, Navigation, neue Fenster, Grimpo

class _Hits:
    """Kleiner Server auf 127.0.0.1: zählt Aufrufe, /start leitet nach Location weiter."""

    def __init__(self, location=""):
        from http.server import BaseHTTPRequestHandler, HTTPServer
        hits = self.hits = []

        class H(BaseHTTPRequestHandler):
            def do_GET(self):
                hits.append(self.path)
                if self.path == "/start" and location:
                    self.send_response(302)
                    self.send_header("Location", location)
                    self.end_headers()
                    return
                body = b"<html><head><title>Intern</title><meta name='citation_doi' content='10.1/x'></head></html>"
                self.send_response(200)
                self.send_header("Content-Type", "text/html")
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)

            def log_message(self, *a):
                pass

        self.srv = HTTPServer(("127.0.0.1", 0), H)
        self.port = self.srv.server_port
        threading.Thread(target=self.srv.serve_forever, daemon=True).start()

    def close(self):
        self.srv.shutdown()
        self.srv.server_close()


class Netz(unittest.TestCase):
    """Links aus Angaben und Fundstücken erreichen nur das öffentliche Internet."""

    def setUp(self):
        self.env = mock.patch.dict(os.environ, {"no_proxy": "*", "NO_PROXY": "*"})
        self.env.start()

    def tearDown(self):
        self.env.stop()

    def test_url_rules(self):
        from benthic import netguard
        for url in ("http://127.0.0.1:8327/api/settings", "http://localhost/", "http://192.168.1.1/",
                    "http://[::1]/", "http://169.254.169.254/latest/meta-data", "http://printer.local/",
                    "file:///etc/passwd", "http://10.0.0.1:99999/"):
            self.assertFalse(netguard.url_allowed(url), url)
        self.assertTrue(netguard.url_allowed("https://www.oecd.org/report"))
        self.assertFalse(netguard.is_public("::ffff:127.0.0.1"))
        self.assertFalse(netguard.is_public("100.64.0.1"))
        self.assertTrue(netguard.is_public("8.8.8.8"))

    def test_local_links_are_not_fetched(self):
        from benthic import clips, verify
        srv = _Hits()
        try:
            self.assertEqual(verify._fetch_page(f"http://127.0.0.1:{srv.port}/"), (None, ""))
            self.assertIsNone(clips.doi_of_page(f"http://localhost:{srv.port}/"))
            with mock.patch.object(verify, "wayback") as wb:
                verify._web_source(f"Seite. http://127.0.0.1:{srv.port}/x", "Seite")
                wb.assert_not_called()  # interne Adressen gehen nicht ans Internet Archive
            self.assertEqual(srv.hits, [])
        finally:
            srv.close()

    def test_names_and_redirects_into_the_local_net_are_blocked(self):
        """Ein Name, der auf 127.0.0.1 zeigt (DNS-Rebinding), und eine Weiterleitung ins eigene Netz."""
        import socket
        import requests
        from benthic import netguard, verify
        real = socket.getaddrinfo

        def resolve(host, port, *args, **kwargs):
            ips = {"outer.example": "127.0.0.1", "inner.example": "10.0.0.5"}
            if host in ips:
                return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", (ips[host], port))]
            return real(host, port, *args, **kwargs)

        srv = _Hits()
        target = _Hits(location=f"http://inner.example:{srv.port}/secret")
        try:
            with mock.patch.object(socket, "getaddrinfo", side_effect=resolve):
                self.assertEqual(verify._fetch_page(f"http://outer.example:{srv.port}/"), (None, ""))
                self.assertEqual(srv.hits, [])
                # Hält der Wächter 127.0.0.1 für öffentlich, folgt er der ersten Antwort, nicht aber der Weiterleitung.
                with mock.patch.object(netguard, "is_public", side_effect=lambda ip: ip == "127.0.0.1"):
                    with netguard.public_only(f"http://outer.example:{target.port}/start") as guard:
                        with self.assertRaises(requests.exceptions.ConnectionError):
                            requests.get(f"http://outer.example:{target.port}/start", timeout=5)
                    self.assertTrue(guard["blocked"])
                self.assertEqual(target.hits, ["/start"])
                self.assertEqual(srv.hits, [])
                # Außerhalb des Blocks verbindet requests wie immer (andere Dienste bleiben unberührt).
                self.assertEqual(requests.get(f"http://outer.example:{srv.port}/ok", timeout=5).status_code, 200)
        finally:
            srv.close()
            target.close()


class Schluessel(unittest.TestCase):
    """API-Keys gehen nur an genau ihren Dienst."""

    def test_keys_by_exact_host(self):
        from benthic import oa
        resp = mock.Mock(status_code=200)
        with mock.patch.object(oa, "_openalex_key", return_value="OAKEY"), \
                mock.patch.object(oa, "_setting", return_value="S2KEY"), \
                mock.patch.object(oa.requests, "get", return_value=resp) as get:
            oa._api_get("https://api.crossref.org/works/10.1/api.openalex.org")
            oa._api_get("https://api.openalex.org.evil.example/works")
            oa._api_get("https://api.unpaywall.org/v2/10.1/api.semanticscholar.org")
            for call in get.call_args_list:
                self.assertNotIn("Authorization", call.kwargs["headers"])
                self.assertNotIn("x-api-key", call.kwargs["headers"])
            get.reset_mock()
            oa._api_get("https://api.openalex.org/works/W1")
            oa._api_get("https://api.semanticscholar.org/graph/v1/paper/x")
            self.assertEqual(get.call_args_list[0].kwargs["headers"]["Authorization"], "Bearer OAKEY")
            self.assertEqual(get.call_args_list[1].kwargs["headers"]["x-api-key"], "S2KEY")

    def test_elsevier_key_stays_with_elsevier(self):
        from benthic import oa
        html = b"<html><head><meta name='citation_pdf_url' content='https://other.example/a.pdf'></head></html>"
        first = mock.Mock(status_code=200, content=html, text=html.decode(), url="https://api.elsevier.com/x",
                          headers={"Content-Type": "text/html"})
        second = mock.Mock(status_code=200, content=b"%PDF-1.4", headers={"Content-Type": "application/pdf"})
        with mock.patch.object(oa.requests, "get", side_effect=[first, second]) as get:
            self.assertEqual(oa.download_pdf("https://api.elsevier.com/x", headers={"X-ELS-APIKey": "ELS"}), b"%PDF-1.4")
        self.assertEqual(get.call_args_list[0].kwargs["headers"]["X-ELS-APIKey"], "ELS")
        self.assertEqual(get.call_args_list[1].args[0], "https://other.example/a.pdf")
        self.assertNotIn("X-ELS-APIKey", get.call_args_list[1].kwargs["headers"])


class Export(unittest.TestCase):
    """Werte aus fremden Metadaten werden in CSV und Excel keine Formeln."""

    def test_safe_cells(self):
        import csv
        import io
        from benthic import batch, export
        self.assertEqual(export.safe_cell("=HYPERLINK(\"https://x\")"), "'=HYPERLINK(\"https://x\")")
        for raw in ("+1", "-1+2", "@SUM(A1)", "\tx", "\rx"):
            self.assertEqual(export.safe_cell(raw), "'" + raw)
        self.assertEqual(export.safe_cell("A study"), "A study")
        self.assertEqual(export.safe_cell(5), 5)
        rows = list(csv.reader(io.StringIO(export.to_csv([{"identifier": "10.1/x", "status": "ok",
                                                           "title": "=cmd|' /C calc'!A0"}]))))
        self.assertIn("'=cmd|' /C calc'!A0", rows[1])
        xlsx = export.to_xlsx([{"identifier": "10.1/x", "status": "ok", "title": "=1+1"}])
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(xlsx)).active
        self.assertIn("'=1+1", [c.value for c in ws[2]])
        b = {"name": "B", "rows": [{"name": "=2+3", "state": "done"}]}
        self.assertIn("'=2+3", batch.to_csv(b, lambda s: s))
        ws = load_workbook(io.BytesIO(batch.to_xlsx(b, lambda s: s))).active
        self.assertIn("'=2+3", [c.value for c in ws[2]])

if __name__ == "__main__":
    unittest.main()
