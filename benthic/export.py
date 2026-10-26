"""Export der Ergebnis-Tabelle als CSV oder XLSX."""

from __future__ import annotations

import csv
import io
from typing import Dict, List, Optional

from .i18n import tr

# Tabellenprogramme lesen Zellen, die so beginnen, als Formel (=HYPERLINK(…), =cmd|…). Werte aus fremden
# Metadaten (Titel, Hinweise, Eingaben) bekommen davor ein Apostroph und bleiben so Text.
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def safe_cell(value):
    if isinstance(value, str) and value.startswith(_FORMULA_START):
        return "'" + value
    return value

COLUMNS = [
    ("identifier", "DOI / Identifier"),
    ("status", "Status"),
    ("source", "Quelle"),
    ("channel", "Kanal"),
    ("legal", "Open Access"),
    ("version", "Fassung"),
    ("filename", "Datei"),
    ("title", "Titel"),
    ("error", "Fehler"),
]

_STATUS_LABEL = {
    "ok": "Geladen",
    "skipped": "Vorhanden",
    "blocked": "Manuell öffnen",
    "present": "Vorhanden",
    "oa": "OA verfügbar",
    "absent": "Nicht gefunden",
    "error": "Fehler",
}

_VERSION_LABEL = {
    "published": "Verlagsversion",
    "accepted": "Akzeptiertes Manuskript",
    "submitted": "Preprint",
}

_CHANNEL_LABEL = {
    "oa": "Open Access",
    "library": "Bibliothek",
}


def _status(r: Dict) -> str:
    status = str(r.get("status") or "")
    # Ältere Läufe speicherten "OA bekannt, PDF blockiert" noch als Fehler.
    if status == "error" and "Open Access bekannt" in str(r.get("error") or ""):
        status = "blocked"
    return tr(_STATUS_LABEL[status]) if status in _STATUS_LABEL else status


def _row(r: Dict) -> Dict[str, str]:
    legal = r.get("legal")
    return {
        "DOI / Identifier": str(r.get("identifier") or ""),
        "Status": _status(r),
        "Quelle": str(r.get("source") or ""),
        "Kanal": (tr(_CHANNEL_LABEL[str(r.get("channel"))]) if str(r.get("channel")) in _CHANNEL_LABEL
                  else str(r.get("route_label") or r.get("channel") or "")),
        "Open Access": tr("ja") if legal is True else (tr("nein") if legal is False else ""),
        "Fassung": tr(_VERSION_LABEL[str(r.get("version"))]) if str(r.get("version")) in _VERSION_LABEL else "",
        "Datei": str(r.get("filename") or ""),
        "Titel": str(r.get("title") or ""),
        "Fehler": str(r.get("error") or ""),
    }


# Prüfbericht (Hallucination-Check): eigene Spalten.
VERIFY_COLUMNS = ["Nr.", "Urteil", "Übereinstimmung", "Angabe: Titel", "Angabe: Jahr",
                  "DOI", "Gefunden: Titel", "Gefunden: Jahr", "Zeitschrift", "Zitiert", "Hinweis", "Eingabe"]
_VERDICT_LABEL = {
    "verified": "Bestätigt", "corrected": "DOI korrigiert", "partial": "Weicht ab",
    "uncertain": "Unsicher", "mismatch": "Passt nicht", "not_found": "Nicht auffindbar",
    "retracted": "Zurückgezogen",
}


def _is_verify(results: List[Dict]) -> bool:
    return bool(results) and "verdict" in results[0]


def _verify_row(r: Dict) -> Dict[str, str]:
    claimed = r.get("claimed") or {}
    found = r.get("found") or {}
    return {
        "Nr.": str((r.get("order") or 0) + 1),
        "Urteil": (tr(_VERDICT_LABEL[str(r.get("verdict"))]) if str(r.get("verdict")) in _VERDICT_LABEL
                   else str(r.get("verdict") or "")),
        "Übereinstimmung": f"{r.get('score', 0)} %",
        "Angabe: Titel": str(claimed.get("title") or ""),
        "Angabe: Jahr": str(claimed.get("year") or ""),
        "DOI": str(r.get("doi") or ""),
        "Gefunden: Titel": str(found.get("title") or ""),
        "Gefunden: Jahr": str(found.get("year") or ""),
        "Zeitschrift": str(found.get("venue") or ""),
        "Zitiert": "" if found.get("cited_by") is None else str(found.get("cited_by")),
        "Hinweis": str(r.get("note") or ""),
        "Eingabe": str(r.get("input") or ""),
    }


def _table(results: List[Dict]):
    """(Spaltennamen, Zeilen) passend zur Art der Ergebnisse."""
    if _is_verify(results):
        rows = sorted(results, key=lambda r: r.get("order", 0))
        return VERIFY_COLUMNS, [_verify_row(r) for r in rows]
    return [c[1] for c in COLUMNS], [_row(r) for r in results]


def table_csv(headers: List[str], rows: List[Dict]) -> str:
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow([tr(h) for h in headers])  # Spaltennamen in der eingestellten Sprache
    for row in rows:
        writer.writerow([safe_cell(row.get(h, "")) for h in headers])
    return out.getvalue()


def table_xlsx(title: str, headers: List[str], rows: List[Dict], widths: Optional[Dict[str, int]] = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = tr(title)[:31]
    ws.append([tr(h) for h in headers])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="08697D")
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append([safe_cell(row.get(h, "")) for h in headers])
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, 22)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


WIDTHS = {
    "DOI / Identifier": 40, "Status": 12, "Quelle": 22, "Kanal": 14,
    "Open Access": 12, "Fassung": 20, "Datei": 48, "Titel": 48, "Fehler": 40,
    "Nr.": 5, "Urteil": 16, "Übereinstimmung": 10, "Angabe: Titel": 48, "Angabe: Jahr": 8,
    "DOI": 32, "Gefunden: Titel": 48, "Gefunden: Jahr": 8, "Zeitschrift": 28, "Zitiert": 8,
    "Hinweis": 48, "Eingabe": 60,
}


def to_csv(results: List[Dict]) -> str:
    return table_csv(*_table(results))


def to_xlsx(results: List[Dict]) -> bytes:
    headers, rows = _table(results)
    return table_xlsx("Prüfbericht" if _is_verify(results) else "Ergebnisse", headers, rows, WIDTHS)


# Quellenmatrix eines Projekts: was Benthic weiß, und die Felder der Lesematrix, die beim Lesen gefüllt werden
# (Forschungsfrage, Theorie, Begriffe, Methode, Kernergebnis, Grenzen, Nutzen für die eigene Frage).
MATRIX_COLUMNS = ["Autor:innen und Jahr", "Titel", "Zeitschrift oder Verlag", "DOI", "Art", "Urteil", "Volltext",
                  "Forschungsfrage der Quelle", "Theorie oder Modell", "Begriffe und Definitionen", "Methode und Stichprobe",
                  "Kernergebnis", "Grenzen", "Nutzen für meine Frage"]
MATRIX_WIDTHS = {"Autor:innen und Jahr": 26, "Titel": 48, "Zeitschrift oder Verlag": 28, "DOI": 30, "Art": 14,
                 "Urteil": 14, "Volltext": 10, "Forschungsfrage der Quelle": 30, "Theorie oder Modell": 26,
                 "Begriffe und Definitionen": 26, "Methode und Stichprobe": 28, "Kernergebnis": 36, "Grenzen": 28,
                 "Nutzen für meine Frage": 32}
_TYPE_LABEL = {"article": "Artikel", "book": "Buch", "book-chapter": "Buchkapitel", "preprint": "Preprint",
               "review": "Übersicht", "dissertation": "Dissertation", "report": "Bericht", "dataset": "Datensatz"}


def matrix_rows(items: List[Dict], meta: Dict[str, Dict]) -> List[Dict]:
    rows = []
    for it in items:
        m = meta.get(str(it.get("doi") or "").lower()) or {}
        authors = m.get("authors") or []
        who = ", ".join(authors[:2]) + (" " + tr("u. a.") if len(authors) > 2 else "")
        year = m.get("year") or it.get("year") or ""
        verdict = str(it.get("verdict") or "")
        rows.append({
            "Autor:innen und Jahr": f"{who} ({year})" if who else str(year),
            "Titel": str(m.get("title") or it.get("title") or ""),
            "Zeitschrift oder Verlag": str(m.get("venue") or ""),
            "DOI": str(it.get("doi") or ""),
            "Art": tr(_TYPE_LABEL[m["type"]]) if m.get("type") in _TYPE_LABEL else str(m.get("type") or ""),
            "Urteil": tr(_VERDICT_LABEL[verdict]) if verdict in _VERDICT_LABEL else "",
            "Volltext": tr("ja") if it.get("in_library") else tr("nein"),
        })
    return rows
